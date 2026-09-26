"""
Excel file generation
"""

import logging
from typing import List, Dict, Any
from pathlib import Path

logger = logging.getLogger(__name__)


class ExcelGenerator:
    """Generate Excel file with Q&A in multiple sheets"""
    
    @staticmethod
    def generate(
        output_path: str,
        english_qa: List[Dict[str, str]],
        hindi_qa: List[Dict[str, str]],
        marathi_qa: List[Dict[str, str]]
    ) -> None:
        """
        Generate Excel file with three sheets
        
        Args:
            output_path: Path to output Excel file
            english_qa: English Q&A pairs
            hindi_qa: Hindi Q&A pairs
            marathi_qa: Marathi Q&A pairs
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl not installed. Install with: pip install openpyxl")
        
        # Create workbook
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default sheet
        
        # Add sheets
        ExcelGenerator._add_sheet(wb, 'English', english_qa)
        ExcelGenerator._add_sheet(wb, 'Hindi', hindi_qa)
        ExcelGenerator._add_sheet(wb, 'Marathi', marathi_qa)
        
        # Save
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        logger.info(f"Excel file saved: {output_path}")
    
    @staticmethod
    def _add_sheet(
        workbook,
        sheet_name: str,
        qa_pairs: List[Dict[str, str]]
    ) -> None:
        """
        Add a sheet to workbook
        
        Args:
            workbook: openpyxl workbook
            sheet_name: Sheet name
            qa_pairs: List of Q&A pairs
        """
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        ws = workbook.create_sheet(sheet_name)
        
        # Set column widths
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 60
        
        # Header styling
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Cell border
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Add headers
        ws['A1'] = 'Questions'
        ws['B1'] = 'Answers'
        
        for cell in ['A1', 'B1']:
            ws[cell].fill = header_fill
            ws[cell].font = header_font
            ws[cell].alignment = header_alignment
            ws[cell].border = thin_border
        
        # Freeze header row
        ws.freeze_panes = 'A2'
        
        # Add data
        for row_idx, qa_pair in enumerate(qa_pairs, start=2):
            question = qa_pair.get('question', '')
            answer = qa_pair.get('answer', '')
            
            ws[f'A{row_idx}'] = question
            ws[f'B{row_idx}'] = answer
            
            # Apply formatting
            for col in ['A', 'B']:
                cell = ws[f'{col}{row_idx}']
                cell.border = thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        
        # Auto-adjust row heights
        ws.row_dimensions[1].height = 25
        for row_idx in range(2, len(qa_pairs) + 2):
            ws.row_dimensions[row_idx].height = None  # Auto height
